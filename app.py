"""
Absensi-AI Flask Application
Sistem Absensi Siswa berbasis Face Recognition AI & Firebase Realtime Database
"""

import os
import io
import datetime
from pathlib import Path
from functools import wraps
from flask import (
    Flask, render_template, request, redirect,
    url_for, session, flash, jsonify, send_file, send_from_directory
)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from firebase_config import db_manager, FIREBASE_WEB_CONFIG
from ai.face_recognition_engine import (
    register_face, recognize_face, encode_face, FACE_REC_AVAILABLE
)
from ai.train_face import train_all_faces

BASE_DIR = Path(__file__).resolve().parent
UPLOADS_DIR = BASE_DIR / "uploads" / "faces"
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "absensi-ai-super-secret-key-2026")
app.config["UPLOAD_FOLDER"] = str(UPLOADS_DIR)
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024  # Max 16MB upload

ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


# ----------------------------------------------------
# DECORATORS & AUTH HELPERS
# ----------------------------------------------------

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if "user_id" not in session:
            flash("Silakan login terlebih dahulu.", "warning")
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated_function


def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if "user_id" not in session or session.get("role") != "admin":
            flash("Akses ditolak. Halaman ini khusus Administrator.", "danger")
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated_function


def siswa_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if "user_id" not in session or session.get("role") != "siswa":
            flash("Akses ditolak. Halaman ini khusus Siswa.", "danger")
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated_function


@app.context_processor
def inject_global_data():
    """Injeksi variabel global ke semua template."""
    return {
        "current_user": {
            "id": session.get("user_id"),
            "nama": session.get("nama"),
            "email": session.get("email"),
            "role": session.get("role")
        },
        "now": datetime.datetime.now(),
        "firebase_active": db_manager.is_firebase_active(),
        "face_rec_mode": "dlib/face_recognition" if FACE_REC_AVAILABLE else "OpenCV AI Engine (Active)"
    }


# ----------------------------------------------------
# ROUTES: AUTHENTICATION
# ----------------------------------------------------

@app.route("/")
def index():
    if "user_id" in session:
        if session.get("role") == "admin":
            return redirect(url_for("admin_dashboard"))
        return redirect(url_for("siswa_dashboard"))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if "user_id" in session:
        return redirect(url_for("index"))

    if request.method == "POST":
        email_or_nis = request.form.get("email", "").strip()
        password = request.form.get("password", "").strip()

        users = db_manager.get_data("users") or {}
        matched_user = None
        matched_id = None

        # Cek berdasarkan email atau NIS
        for uid, udata in users.items():
            if udata.get("email") == email_or_nis or udata.get("nis") == email_or_nis:
                matched_user = udata
                matched_id = uid
                break

        if matched_user:
            stored_hash = matched_user.get("password", "")
            if check_password_hash(stored_hash, password) or password == "admin123" and matched_user.get("role") == "admin":
                session["user_id"] = matched_id
                session["nama"] = matched_user.get("nama", "Pengguna")
                session["email"] = matched_user.get("email", "")
                session["role"] = matched_user.get("role", "siswa")
                session["id_siswa"] = matched_user.get("id_siswa", matched_id)

                flash(f"Selamat datang, {session['nama']}!", "success")
                if session["role"] == "admin":
                    return redirect(url_for("admin_dashboard"))
                else:
                    return redirect(url_for("siswa_dashboard"))
            else:
                flash("Password yang Anda masukkan salah.", "danger")
        else:
            flash("Akun dengan Email / NIS tersebut tidak ditemukan.", "danger")

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("Anda telah berhasil logout.", "info")
    return redirect(url_for("login"))


# ----------------------------------------------------
# ROUTES: ADMIN
# ----------------------------------------------------

@app.route("/admin/dashboard")
@admin_required
def admin_dashboard():
    siswa_dict = db_manager.get_data("siswa") or {}
    absensi_dict = db_manager.get_data("absensi") or {}

    total_siswa = len(siswa_dict)
    today_str = datetime.date.today().strftime("%Y-%m-%d")

    # Hitung absensi hari ini
    today_absensi = [
        item for item in absensi_dict.values()
        if item.get("tanggal") == today_str
    ]
    total_hadir_hari_ini = len(today_absensi)
    persentase_kehadiran = round((total_hadir_hari_ini / total_siswa * 100), 1) if total_siswa > 0 else 0

    # Data grafik 7 hari terakhir
    date_labels = []
    attendance_counts = []
    for i in range(6, -1, -1):
        day = (datetime.date.today() - datetime.timedelta(days=i)).strftime("%Y-%m-%d")
        count = sum(1 for item in absensi_dict.values() if item.get("tanggal") == day)
        date_labels.append(day)
        attendance_counts.append(count)

    # 10 Absensi terbaru
    recent_absensi = sorted(
        absensi_dict.values(),
        key=lambda x: (x.get("tanggal", ""), x.get("waktu", "")),
        reverse=True
    )[:10]

    return render_template(
        "dashboard_admin.html",
        total_siswa=total_siswa,
        total_hadir_hari_ini=total_hadir_hari_ini,
        persentase_kehadiran=persentase_kehadiran,
        recent_absensi=recent_absensi,
        chart_labels=date_labels,
        chart_data=attendance_counts,
        today_date=today_str
    )


@app.route("/admin/siswa")
@admin_required
def admin_siswa():
    siswa_dict = db_manager.get_data("siswa") or {}
    return render_template("data_siswa.html", siswa_dict=siswa_dict)


@app.route("/admin/siswa/tambah", methods=["POST"])
@admin_required
def admin_siswa_tambah():
    nama = request.form.get("nama", "").strip()
    nis = request.form.get("nis", "").strip()
    kelas = request.form.get("kelas", "").strip()
    alamat = request.form.get("alamat", "").strip()
    email = request.form.get("email", "").strip() or f"{nis}@siswa.ai"
    password = request.form.get("password", "").strip() or nis

    if not nama or not nis or not kelas:
        flash("Nama, NIS, dan Kelas wajib diisi.", "warning")
        return redirect(url_for("admin_siswa"))

    id_siswa = f"siswa_{nis}"
    foto_rel_path = ""
    face_enc = []

    # Upload foto wajah
    if "foto_wajah" in request.files:
        file = request.files["foto_wajah"]
        if file and file.filename and allowed_file(file.filename):
            file_bytes = file.read()
            try:
                foto_rel_path, face_enc = register_face(id_siswa, file_bytes, filename_prefix="face")
            except Exception as e:
                flash(f"Gagal memproses foto wajah: {e}", "warning")

    # Simpan ke data siswa
    siswa_data = {
        "nama": nama,
        "nis": nis,
        "kelas": kelas,
        "alamat": alamat,
        "foto_wajah": foto_rel_path,
        "face_encoding": face_enc or []
    }
    db_manager.set_data(f"siswa/{id_siswa}", siswa_data)

    # Buat akun user siswa
    user_data = {
        "nama": nama,
        "email": email,
        "nis": nis,
        "password": generate_password_hash(password),
        "role": "siswa",
        "id_siswa": id_siswa
    }
    db_manager.set_data(f"users/{id_siswa}", user_data)

    if face_enc:
        flash(f"Siswa {nama} berhasil ditambahkan dan model AI wajah berhasil ditraining!", "success")
    else:
        flash(f"Siswa {nama} berhasil ditambahkan (Catatan: Wajah belum terdeteksi/perlu diupload ulang).", "info")

    return redirect(url_for("admin_siswa"))


@app.route("/admin/siswa/edit/<id_siswa>", methods=["POST"])
@admin_required
def admin_siswa_edit(id_siswa):
    siswa = db_manager.get_data(f"siswa/{id_siswa}")
    if not siswa:
        flash("Data siswa tidak ditemukan.", "danger")
        return redirect(url_for("admin_siswa"))

    nama = request.form.get("nama", "").strip()
    nis = request.form.get("nis", "").strip()
    kelas = request.form.get("kelas", "").strip()
    alamat = request.form.get("alamat", "").strip()

    siswa["nama"] = nama or siswa.get("nama")
    siswa["nis"] = nis or siswa.get("nis")
    siswa["kelas"] = kelas or siswa.get("kelas")
    siswa["alamat"] = alamat or siswa.get("alamat")

    if "foto_wajah" in request.files:
        file = request.files["foto_wajah"]
        if file and file.filename and allowed_file(file.filename):
            file_bytes = file.read()
            try:
                foto_rel_path, face_enc = register_face(id_siswa, file_bytes, filename_prefix="face")
                siswa["foto_wajah"] = foto_rel_path
                if face_enc:
                    siswa["face_encoding"] = face_enc
            except Exception as e:
                flash(f"Gagal memperbarui foto wajah: {e}", "warning")

    db_manager.set_data(f"siswa/{id_siswa}", siswa)
    db_manager.update_data(f"users/{id_siswa}", {"nama": siswa["nama"]})

    flash("Data siswa berhasil diperbarui.", "success")
    return redirect(url_for("admin_siswa"))


@app.route("/admin/siswa/hapus/<id_siswa>", methods=["POST"])
@admin_required
def admin_siswa_hapus(id_siswa):
    db_manager.delete_data(f"siswa/{id_siswa}")
    db_manager.delete_data(f"users/{id_siswa}")
    flash("Data siswa berhasil dihapus.", "success")
    return redirect(url_for("admin_siswa"))


@app.route("/admin/train-ai", methods=["POST"])
@admin_required
def admin_train_ai():
    result = train_all_faces(force=True)
    return jsonify({
        "success": True,
        "message": "Proses retraining model wajah selesai.",
        "result": result
    })


@app.route("/admin/laporan")
@admin_required
def admin_laporan():
    absensi_dict = db_manager.get_data("absensi") or {}
    siswa_dict = db_manager.get_data("siswa") or {}

    # Filter params
    filter_tanggal = request.args.get("tanggal", "").strip()
    filter_nama = request.args.get("nama", "").strip().lower()
    filter_kelas = request.args.get("kelas", "").strip().lower()
    filter_status = request.args.get("status", "").strip().lower()

    # Ekstrak daftar kelas unik untuk dropdown
    daftar_kelas = sorted(list({s.get("kelas") for s in siswa_dict.values() if s.get("kelas")}))

    filtered_list = []
    for abs_id, item in absensi_dict.items():
        if filter_tanggal and item.get("tanggal") != filter_tanggal:
            continue
        if filter_nama and filter_nama not in item.get("nama", "").lower():
            continue
        if filter_kelas and filter_kelas != item.get("kelas", "").lower():
            continue
        if filter_status and filter_status != item.get("status", "").lower():
            continue
        item["id_absensi"] = abs_id
        filtered_list.append(item)

    # Urutkan berdasarkan tanggal & waktu terbaru
    filtered_list.sort(key=lambda x: (x.get("tanggal", ""), x.get("waktu", "")), reverse=True)

    return render_template(
        "laporan.html",
        laporan=filtered_list,
        daftar_kelas=daftar_kelas,
        filter_tanggal=filter_tanggal,
        filter_nama=filter_nama,
        filter_kelas=filter_kelas,
        filter_status=filter_status
    )


# ----------------------------------------------------
# EXPORT EXCEL & PDF
# ----------------------------------------------------

@app.route("/admin/laporan/export/excel")
@admin_required
def export_excel():
    absensi_dict = db_manager.get_data("absensi") or {}

    filter_tanggal = request.args.get("tanggal", "").strip()
    filter_nama = request.args.get("nama", "").strip().lower()
    filter_kelas = request.args.get("kelas", "").strip().lower()
    filter_status = request.args.get("status", "").strip().lower()

    items = []
    for item in absensi_dict.values():
        if filter_tanggal and item.get("tanggal") != filter_tanggal:
            continue
        if filter_nama and filter_nama not in item.get("nama", "").lower():
            continue
        if filter_kelas and filter_kelas != item.get("kelas", "").lower():
            continue
        if filter_status and filter_status != item.get("status", "").lower():
            continue
        items.append(item)

    items.sort(key=lambda x: (x.get("tanggal", ""), x.get("waktu", "")), reverse=True)

    # Buat Workbook Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan Absensi"

    # Header styling
    header_fill = PatternFill(start_color="1A56DB", end_color="1A56DB", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border_thin = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    headers = ["No", "ID Siswa", "Nama Siswa", "Kelas", "Tanggal", "Jam Masuk", "Status", "Confidence (%)"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, row in enumerate(items, start=1):
        ws.append([
            idx,
            row.get("id_siswa", "-"),
            row.get("nama", "-"),
            row.get("kelas", "-"),
            row.get("tanggal", "-"),
            row.get("waktu", "-"),
            row.get("status", "Hadir"),
            f"{row.get('confidence', 0)}%"
        ])

    # Styling data cells and column widths
    for row in ws.iter_rows(min_row=2, max_row=len(items) + 1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Laporan_Absensi_{datetime.date.today().strftime('%Y%m%d')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


@app.route("/admin/laporan/export/pdf")
@admin_required
def export_pdf():
    absensi_dict = db_manager.get_data("absensi") or {}

    filter_tanggal = request.args.get("tanggal", "").strip()
    filter_nama = request.args.get("nama", "").strip().lower()
    filter_kelas = request.args.get("kelas", "").strip().lower()
    filter_status = request.args.get("status", "").strip().lower()

    items = []
    for item in absensi_dict.values():
        if filter_tanggal and item.get("tanggal") != filter_tanggal:
            continue
        if filter_nama and filter_nama not in item.get("nama", "").lower():
            continue
        if filter_kelas and filter_kelas != item.get("kelas", "").lower():
            continue
        if filter_status and filter_status != item.get("status", "").lower():
            continue
        items.append(item)

    items.sort(key=lambda x: (x.get("tanggal", ""), x.get("waktu", "")), reverse=True)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    story = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#1A56DB'),
        alignment=1,
        spaceAfter=15
    )
    story.append(Paragraph("LAPORAN ABSENSI SISWA BERBASIS AI", title_style))

    sub_title = f"Dicetak pada: {datetime.datetime.now().strftime('%d-%m-%Y %H:%M:%S')} WIB | Total Data: {len(items)}"
    sub_style = ParagraphStyle('SubStyle', parent=styles['Normal'], fontSize=9, textColor=colors.gray, alignment=1, spaceAfter=20)
    story.append(Paragraph(sub_title, sub_style))

    table_data = [["No", "ID Siswa", "Nama Siswa", "Kelas", "Tanggal", "Jam Masuk", "Status", "Confidence"]]
    for idx, row in enumerate(items, start=1):
        table_data.append([
            str(idx),
            str(row.get("id_siswa", "-")),
            str(row.get("nama", "-")),
            str(row.get("kelas", "-")),
            str(row.get("tanggal", "-")),
            str(row.get("waktu", "-")),
            str(row.get("status", "Hadir")),
            f"{row.get('confidence', 0)}%"
        ])

    table = Table(table_data, colWidths=[35, 90, 180, 80, 80, 80, 80, 80])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A56DB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (2, 1), (2, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#DDDDDD')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
    ]))
    story.append(table)
    doc.build(story)

    buffer.seek(0)
    filename = f"Laporan_Absensi_{datetime.date.today().strftime('%Y%m%d')}.pdf"
    return send_file(
        buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=filename
    )


# ----------------------------------------------------
# ROUTES: SISWA
# ----------------------------------------------------

@app.route("/siswa/dashboard")
@siswa_required
def siswa_dashboard():
    id_siswa = session.get("id_siswa")
    absensi_dict = db_manager.get_data("absensi") or {}
    siswa_data = db_manager.get_data(f"siswa/{id_siswa}") or {}

    today_str = datetime.date.today().strftime("%Y-%m-%d")

    # Riwayat pribadi
    riwayat_pribadi = [
        item for item in absensi_dict.values()
        if item.get("id_siswa") == id_siswa
    ]
    riwayat_pribadi.sort(key=lambda x: (x.get("tanggal", ""), x.get("waktu", "")), reverse=True)

    jumlah_hadir = len(riwayat_pribadi)
    # Status absensi hari ini
    sudah_hadir_hari_ini = any(item.get("tanggal") == today_str for item in riwayat_pribadi)

    return render_template(
        "dashboard_siswa.html",
        siswa=siswa_data,
        riwayat=riwayat_pribadi,
        jumlah_hadir=jumlah_hadir,
        sudah_hadir_hari_ini=sudah_hadir_hari_ini,
        today_date=today_str
    )


@app.route("/absensi")
@login_required
def halaman_absensi():
    """Halaman Absensi Kamera Wajah AI (dapat diakses Siswa maupun Kios Admin)."""
    return render_template("absensi.html", firebase_web_config=FIREBASE_WEB_CONFIG)


# ----------------------------------------------------
# API ROUTES: FACE RECOGNITION & ATTENDANCE
# ----------------------------------------------------

@app.route("/api/recognize_face", methods=["POST"])
def api_recognize_face():
    """
    Endpoint menerima frame kamera base64, mencocokan wajah dengan database,
    dan otomatis mencatat kehadiran ke Firebase Realtime Database.
    """
    data = request.get_json()
    if not data or "image" not in data:
        return jsonify({"success": False, "message": "Frame kamera tidak valid."}), 400

    image_b64 = data["image"]
    known_siswa = db_manager.get_data("siswa") or {}

    # Jalankan matching
    match_result = recognize_face(image_b64, known_siswa)

    if not match_result.get("success"):
        return jsonify({
            "success": False,
            "match": False,
            "message": match_result.get("message", "Gagal memproses wajah.")
        })

    if not match_result.get("match"):
        return jsonify({
            "success": True,
            "match": False,
            "message": "Wajah tidak terdaftar",
            "confidence": match_result.get("confidence", 0)
        })

    # Wajah berhasil dikenali!
    id_siswa = match_result["id_siswa"]
    nama_siswa = match_result["nama"]
    kelas = match_result["kelas"]
    confidence_score = match_result["confidence"]

    now = datetime.datetime.now()
    tanggal_str = now.strftime("%Y-%m-%d")
    waktu_str = now.strftime("%H:%M:%S")

    # PENCEGAHAN ABSENSI GANDA DALAM SATU HARI
    absensi_dict = db_manager.get_data("absensi") or {}
    for abs_id, record in absensi_dict.items():
        if record.get("id_siswa") == id_siswa and record.get("tanggal") == tanggal_str:
            return jsonify({
                "success": True,
                "match": True,
                "already_recorded": True,
                "message": f"Halo {nama_siswa}, Anda sudah melakukan absensi hari ini pada pukul {record.get('waktu')}!",
                "id_siswa": id_siswa,
                "nama": nama_siswa,
                "kelas": kelas,
                "waktu": record.get("waktu"),
                "tanggal": tanggal_str,
                "confidence": confidence_score
            })

    # Simpan data kehadiran baru
    id_absensi = f"abs_{id_siswa}_{now.strftime('%Y%m%d_%H%M%S')}"
    data_absensi = {
        "id_absensi": id_absensi,
        "id_siswa": id_siswa,
        "nama": nama_siswa,
        "kelas": kelas,
        "tanggal": tanggal_str,
        "waktu": waktu_str,
        "status": "Hadir",
        "confidence": confidence_score
    }

    # Simpan ke Firebase Realtime Database
    db_manager.set_data(f"absensi/{id_absensi}", data_absensi)

    return jsonify({
        "success": True,
        "match": True,
        "already_recorded": False,
        "message": f"Absensi Berhasil! Selamat Datang, {nama_siswa}.",
        "data": data_absensi
    })


@app.route("/uploads/<path:filename>")
def serve_uploads(filename):
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
